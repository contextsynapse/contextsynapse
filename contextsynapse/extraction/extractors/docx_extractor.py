"""
DOCX Extractor

Extracts content from DOCX files.
"""

from pathlib import Path
from typing import Dict, Any, List
import logging

from ..registry import BaseExtractor, ExtractionResult

logger = logging.getLogger(__name__)


class DocxExtractor(BaseExtractor):
    """
    DOCX extractor using python-docx.
    """
    
    def __init__(self, config):
        """Initialize DOCX extractor."""
        super().__init__(config)
        self.extractor_version = "docx_v1"
    
    def supports_format(self, file_path: str) -> bool:
        """Check if file is a DOCX."""
        return Path(file_path).suffix.lower() == ".docx"
    
    def extract(self, file_path: str) -> ExtractionResult:
        """
        Extract content from DOCX.
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            ExtractionResult
        """
        result = ExtractionResult(extractor_version=self.extractor_version)
        
        try:
            from docx import Document
            
            doc = Document(file_path)
            
            # Metadata
            core_props = doc.core_properties
            result.metadata = {
                "title": core_props.title or "",
                "author": core_props.author or "",
                "subject": core_props.subject or "",
                "created": str(core_props.created) if core_props.created else "",
            }
            
            # Extract paragraphs
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            
            # Extract tables
            tables = []
            for table_idx, table in enumerate(doc.tables, 1):
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                
                if table_data:
                    result.tables.append({
                        "id": f"tbl_{table_idx}",
                        "data": table_data,
                        "headers": table_data[0] if table_data else [],
                        "rows": table_data[1:] if len(table_data) > 1 else [],
                    })
            
            # Combine all paragraphs into a single page
            full_text = "\n\n".join(paragraphs)
            
            result.pages.append({
                "page_no": 1,
                "text": full_text,
                "content": full_text,
            })
            
            result.metadata["page_count"] = 1
        
        except ImportError:
            error_msg = "python-docx not available. Install with: pip install python-docx"
            logger.error(error_msg)
            result.errors.append(error_msg)
        except Exception as e:
            logger.error(f"DOCX extraction failed: {e}")
            result.errors.append(str(e))
        
        return result
































